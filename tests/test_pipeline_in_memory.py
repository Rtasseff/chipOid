"""Integration test for the in-memory pipeline entry point.

Exercises `chipoid.pipeline.run_batch_in_memory` against the seeded
`data/mcf7_media.tif` (+ companions) using an in-memory manifest DataFrame.
This is the canary for the refactor in pipeline.py: GUI invocation path
must produce the same artifacts the CLI does.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from chipoid.config import load_config
from chipoid.pipeline import run_batch_in_memory


# All tests in this file require the seeded data on disk.
DATA_ROOT = Path("data")
SEED_BF = DATA_ROOT / "mcf7_media.tif"
SEED_GREEN = DATA_ROOT / "mcf7_media_green.tif"
SEED_RED = DATA_ROOT / "mcf7_media_red.tif"


pytestmark = pytest.mark.skipif(
    not (SEED_BF.exists() and SEED_GREEN.exists() and SEED_RED.exists()),
    reason="seeded mcf7_media.tif + companions not present under data/",
)


def test_run_batch_in_memory_produces_outputs(tmp_path):
    cfg = load_config(None)  # pure defaults — exhaustive validation
    cfg["input"]["data_root"] = str(DATA_ROOT)
    cfg["output"]["dir"] = str(tmp_path)

    manifest = pd.DataFrame([{
        "image_id": "mcf7_media",
        "source": "mcf7_media.tif",
        # An extra metadata column to confirm propagation into wells_all.csv.
        "condition": "control_test",
    }])

    log_lines: list[str] = []
    result = run_batch_in_memory(cfg, manifest, log=log_lines.append)

    assert result["n_images"] == 1
    assert result["n_success"] == 1
    assert result["n_failed"] == 0

    # Per-image outputs
    per_image_dir = tmp_path / "mcf7_media"
    assert per_image_dir.is_dir()
    assert (per_image_dir / "wells.csv").exists()

    # Consolidated CSV (with manifest metadata propagated)
    wells_all = tmp_path / "wells_all.csv"
    assert wells_all.exists()
    df = pd.read_csv(wells_all)
    assert "condition" in df.columns
    assert (df["condition"] == "control_test").all()
    # Sanity-check the canonical 100-well result from the seed image.
    assert len(df) == 100

    # Batch summary
    assert (tmp_path / "batch_summary.csv").exists()

    # Log callback was actually exercised
    assert len(log_lines) > 0
    assert any("mcf7_media" in line for line in log_lines)


# --------------------------------------------------------------------------- #
# Inclusion + workbook
# --------------------------------------------------------------------------- #
def _run(tmp_path, inclusion=None, **out_over):
    cfg = load_config(None)
    cfg["input"]["data_root"] = str(DATA_ROOT)
    cfg["output"]["dir"] = str(tmp_path)
    cfg["output"].update(out_over)
    if inclusion is not None:
        cfg["inclusion"].update(inclusion)
    manifest = pd.DataFrame([{"image_id": "mcf7_media", "source": "mcf7_media.tif"}])
    logs: list[str] = []
    result = run_batch_in_memory(cfg, manifest, log=logs.append)
    return result, logs


def test_inclusion_off_adds_nothing(tmp_path):
    from openpyxl import load_workbook
    result, logs = _run(tmp_path)
    assert result["n_failed"] == 0
    df = pd.read_csv(tmp_path / "wells_all.csv")
    assert "included" not in df.columns and "exclude_reason" not in df.columns
    assert "n_included" not in pd.read_csv(tmp_path / "batch_summary.csv").columns
    assert load_workbook(tmp_path / "wells_all.xlsx").sheetnames == [
        "all_wells", "summary", "settings"]
    assert logs[0].startswith("chipOid ") and " batch: 1 images" in logs[0]
    assert "inclusion: off — all wells included" in logs
    assert not any(l.strip().startswith("inclusion:") and "wells included (" in l for l in logs)


def test_inclusion_on_columns_summary_workbook(tmp_path):
    from openpyxl import load_workbook
    result, logs = _run(tmp_path, inclusion={"enabled": True,
                                             "min_signal": {"green": 2500, "red": 600}})
    assert result["n_failed"] == 0
    df = pd.read_csv(tmp_path / "wells_all.csv")
    cols = list(df.columns)
    assert cols[cols.index("dist_to_det") + 1: cols.index("dist_to_det") + 3] == [
        "included", "exclude_reason"]
    summ = pd.read_csv(tmp_path / "batch_summary.csv")
    n_inc = int(summ["n_included"].iloc[0])
    assert n_inc == int(df["included"].sum()) and 0 < n_inc < len(df)
    assert summ["min_signal_green"].iloc[0] == 2500 and summ["min_signal_red"].iloc[0] == 600
    wb = load_workbook(tmp_path / "wells_all.xlsx")
    assert wb.sheetnames == ["all_wells", "included_wells", "summary", "settings"]
    assert pd.read_excel(tmp_path / "wells_all.xlsx", sheet_name="included_wells").shape[0] == n_inc
    assert any(l.startswith("inclusion: on — ") for l in logs)
    assert any(l.strip().startswith("inclusion: ") and f"{n_inc}/{len(df)} wells included" in l
               for l in logs)


@pytest.mark.parametrize("inclusion_on", [True, False])
def test_runs_without_signal_in_readout_metrics(tmp_path, inclusion_on):
    """Inclusion and the 06/07 diagnostics use the measurement arrays, so
    neither needs `signal` listed in readout.metrics."""
    cfg = load_config(None)
    cfg["input"]["data_root"] = str(DATA_ROOT)
    cfg["output"]["dir"] = str(tmp_path)
    cfg["readout"]["metrics"] = ["mean"]
    cfg["inclusion"]["enabled"] = inclusion_on
    manifest = pd.DataFrame([{"image_id": "mcf7_media", "source": "mcf7_media.tif"}])
    result = run_batch_in_memory(cfg, manifest, log=lambda m: None)
    assert result["n_failed"] == 0
    assert ("included" in pd.read_csv(tmp_path / "wells_all.csv").columns) == inclusion_on
    assert (tmp_path / "mcf7_media" / "06_histograms.png").exists()
    assert (tmp_path / "mcf7_media" / "07_scatter.png").exists()


def test_shape_mismatch_fails_only_that_image(tmp_path):
    import shutil
    import numpy as np
    import tifffile
    data = tmp_path / "data"; data.mkdir()
    for stem in ("good", "bad"):
        shutil.copy(SEED_BF, data / f"{stem}.tif")
        shutil.copy(SEED_GREEN, data / f"{stem}_green.tif")
        shutil.copy(SEED_RED, data / f"{stem}_red.tif")
    tifffile.imwrite(data / "bad_red.tif", np.zeros((10, 10), dtype=np.uint16))
    cfg = load_config(None)
    cfg["input"]["data_root"] = str(data)
    cfg["output"]["dir"] = str(tmp_path / "out")
    manifest = pd.DataFrame([{"image_id": "bad", "source": "bad.tif"},
                             {"image_id": "good", "source": "good.tif"}])
    logs: list[str] = []
    result = run_batch_in_memory(cfg, manifest, log=logs.append)
    assert result["n_failed"] == 1 and result["n_success"] == 1
    assert any("[bad] shape mismatch for marker 'red'" in l for l in logs)
    assert set(pd.read_csv(tmp_path / "out" / "wells_all.csv").image_id) == {"good"}


def test_missing_companion_fails_image_when_inclusion_on(tmp_path):
    data = tmp_path / "data"; data.mkdir()
    import shutil
    shutil.copy(SEED_BF, data / "x.tif"); shutil.copy(SEED_GREEN, data / "x_green.tif")
    cfg = load_config(None)
    cfg["input"]["data_root"] = str(data)
    cfg["output"]["dir"] = str(tmp_path / "out")
    cfg["inclusion"]["enabled"] = True
    manifest = pd.DataFrame([{"image_id": "x", "source": "x.tif"}])
    logs: list[str] = []
    result = run_batch_in_memory(cfg, manifest, log=logs.append)
    assert result["n_failed"] == 1 and result["n_success"] == 0
    assert any("[x] inclusion is enabled but the 'red' companion is missing" in l for l in logs)


def test_open_files_warn_and_continue(tmp_path, monkeypatch):
    """PermissionError on any output write is a warning, never a failed batch."""
    real_to_csv = pd.DataFrame.to_csv

    def locked(self, path_or_buf=None, *a, **k):
        if str(path_or_buf).endswith(("wells_all.csv", "batch_summary.csv")):
            raise PermissionError(13, "Permission denied", str(path_or_buf))
        return real_to_csv(self, path_or_buf, *a, **k)

    monkeypatch.setattr(pd.DataFrame, "to_csv", locked)
    result, logs = _run(tmp_path)
    assert result["n_failed"] == 0 and result["n_success"] == 1
    assert (tmp_path / "mcf7_media" / "wells.csv").exists()
    assert (tmp_path / "wells_all.xlsx").exists()
    warns = [l for l in logs if "[WARN] could not write" in l]
    assert any("wells_all.csv" in l for l in warns) and any("batch_summary.csv" in l for l in warns)
